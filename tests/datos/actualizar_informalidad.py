"""
Regenera ``dane_informalidad.csv`` con la serie mensual de ocupados formales e informales del DANE, y
``dane_informalidad_marco2005.csv`` con la de la medición anterior.

    uv run --extra app --with xlrd python tests/datos/actualizar_informalidad.py jun-ago2026

El argumento es el anexo vigente de «Empleo informal y seguridad social»: el periodo como aparece en su
nombre (``anex-GEIHEISS-jun-ago2026.xlsx``), una URL o un archivo ya descargado. Está en
https://www.dane.gov.co/index.php/estadisticas-por-tema-2/mercado-laboral/empleo-informal-y-seguridad-social

Se toma el total nacional de la hoja «Grandes dominios» desde 2022: los microdatos de 2021 en el ANDA son
del cuestionario anterior y no tienen las preguntas de la medición.

La medición anterior (marco 2005) se publicó solo para 13 y 23 ciudades, en trimestre móvil. Sale del anexo
de octubre a diciembre de 2021 (``ANEXO_2005``, fijo), el último con las proyecciones del censo 2005, como
los microdatos de 2007–2019 y 2021 en el ANDA. Es un .xls: hace falta ``xlrd``.
"""

from __future__ import annotations

import argparse
import tempfile
from pathlib import Path

import openpyxl
import pandas as pd

from actualizar_publicado import MESES, abrir

ANEXO_VIGENTE = "https://www.dane.gov.co/files/operaciones/GEIH/anex-GEIHEISS-{periodo}.xlsx"
ANEXO_2005 = ("https://www.dane.gov.co/files/investigaciones/boletines/ech/ech_informalidad/"
              "anexo_geih_informalidad_oct21_dic21.xls")
SALIDA = Path(__file__).with_name("dane_informalidad.csv")
SALIDA_2005 = Path(__file__).with_name("dane_informalidad_marco2005.csv")
FILAS = {"Población ocupada": "ocupados", "Formal": "formales", "Informal": "informales"}
# Bloque de la hoja «Ciudades» -> dominio del CSV
DOMINIOS_2005 = {"Total 13 ciudades y AM": "13 ciudades", "23 ciudades y áreas": "23 ciudades"}


def leer(archivo: Path) -> pd.DataFrame:
    """Serie mensual del bloque «Total nacional» de la hoja «Grandes dominios»."""
    libro = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    hoja = next(h for h in libro.sheetnames if h.strip() == "Grandes dominios")
    filas = list(libro[hoja].iter_rows(values_only=True))
    b = next(k for k, r in enumerate(filas) if str(r[0]).strip() == "Total nacional")
    i = next(k for k in range(b, len(filas)) if filas[k][0] == "Concepto")
    anios, anio = [], None
    for v in filas[i][1:]:
        anio = int(str(v)[:4]) if v is not None and str(v).strip()[:4].isdigit() else anio
        anios.append(anio)
    meses = [MESES.get(str(v).strip()[:3]) for v in filas[i + 1][1:]]
    datos = {}
    for r in filas[i + 2:]:
        if r[0] is None:
            break  # fin del bloque
        if (clave := FILAS.get(str(r[0]).strip())) and clave not in datos:
            datos[clave] = [v if isinstance(v, (int, float)) else None for v in r[1:len(anios) + 1]]
    df = pd.DataFrame({"anio": anios, "mes": meses, **datos})
    return df.dropna(subset=["anio", "mes", "ocupados"]).astype({"anio": int, "mes": int})


def leer_marco2005(archivo: Path) -> pd.DataFrame:
    """Trimestres móviles de las 13 y 23 ciudades (hoja «Ciudades»): la columna j termina en marzo de 2007 + j."""
    import xlrd

    hoja = xlrd.open_workbook(archivo).sheet_by_name("Ciudades")
    filas = [[c.value for c in hoja.row(r)] for r in range(hoja.nrows)]
    datos = []
    for r, fila in enumerate(filas[:-4]):
        if (dominio := DOMINIOS_2005.get(str(fila[0]).strip())) is None:
            continue
        for k, concepto in enumerate(("ocupados", "formales", "informales")):
            for j, v in enumerate(filas[r + 2 + k][1:]):
                if isinstance(v, float):
                    fin = pd.Period("2007-03", "M") + j
                    datos.append({"dominio": dominio, "anio": fin.year, "mes": fin.month, "concepto": concepto,
                                  "valor": v})
    return (pd.DataFrame(datos).pivot_table(index=["dominio", "anio", "mes"], columns="concepto", values="valor")
            .reset_index()[["dominio", "anio", "mes", "ocupados", "formales", "informales"]])


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__.strip().splitlines()[0])
    p.add_argument("anexo", help="anexo vigente: periodo (p. ej. jun-ago2026), URL o archivo .xlsx")
    p.add_argument("--salida", type=Path, default=SALIDA)
    args = p.parse_args()

    fuente = args.anexo if args.anexo.startswith("http") or Path(args.anexo).is_file() else \
        ANEXO_VIGENTE.format(periodo=args.anexo)
    with tempfile.TemporaryDirectory() as tmp:
        archivo = abrir(fuente, Path(tmp))
        out = leer(archivo)
    nombre = archivo.stem.replace("anex-GEIHEISS-", "")
    out = out[out.anio >= 2022].assign(anexo=nombre)
    encabezado = f"""\
# Ocupados formales e informales publicados por el DANE (medición de la 17.ª CIET): serie mensual del total
# nacional, en miles. Se regenera con tests/datos/actualizar_informalidad.py.
# Anexo de «Empleo informal y seguridad social», hoja "Grandes dominios":
#   {ANEXO_VIGENTE.format(periodo=nombre)}
# Desde 2022: los microdatos de 2021 en el ANDA no tienen las preguntas de la medición.
"""
    with open(args.salida, "w", encoding="utf-8") as fh:
        fh.write(encabezado)
        out.to_csv(fh, index=False, float_format="%.6f")
    print(f"{len(out)} meses en {args.salida} ({out.anio.min()}-{out.mes.iloc[0]:02d} a "
          f"{out.anio.max()}-{out.mes.iloc[-1]:02d})")

    with tempfile.TemporaryDirectory() as tmp:
        viejo = leer_marco2005(abrir(ANEXO_2005, Path(tmp)))
    encabezado = f"""\
# Ocupados formales e informales con la medición anterior del DANE (marco 2005, 15.ª CIET): trimestre móvil
# (mes = el último del trimestre) de 13 y 23 ciudades y áreas metropolitanas, en miles, con las proyecciones
# del censo 2005. Se regenera con tests/datos/actualizar_informalidad.py. Hoja "Ciudades" de:
#   {ANEXO_2005}
"""
    with open(SALIDA_2005, "w", encoding="utf-8") as fh:
        fh.write(encabezado)
        viejo.to_csv(fh, index=False, float_format="%.6f")
    print(f"{len(viejo)} trimestres en {SALIDA_2005}")


if __name__ == "__main__":
    main()
